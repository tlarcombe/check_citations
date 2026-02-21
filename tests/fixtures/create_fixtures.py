"""Run once to generate binary test fixtures."""
from pathlib import Path
from docx import Document

def create_sample_docx():
    doc = Document()
    doc.add_heading("Test Document", 0)
    doc.add_paragraph("This is a claim about climate change [Smith 2023].")
    doc.add_paragraph("See also https://en.wikipedia.org/wiki/Climate_change")
    out = Path(__file__).parent / "sample.docx"
    doc.save(str(out))
    print(f"Created {out}")

if __name__ == "__main__":
    create_sample_docx()
